"""An in-memory workbook shaped like the owner's sheet, shared by the parser and writer tests."""

from datetime import time

import openpyxl

PACK_HEADER = ["類別", "項目", None, "數量", "已打包數量", "打包狀態", "Double Check",
               "打包時機", "需求", "取得地點", "備註"]
TRANSPORT_HEADER = ["目標起點", "目標終點", "交通工具", "提前買票", "路線圖", "時刻表", "即時動態",
                    "方向", "起點", "終點", "實際起點", "實際終點", "價錢", "時間", "班次間隔",
                    "平日班次 (早)", "平日班次 (中)", "平日班次 (下午)", "平日班次 (晚)",
                    "假日班次 (早)", "假日班次 (中)", "假日班次 (下午)", "假日班次 (晚)"]
THIS_TIME_HEADER = ["出發地點", "目的地", "時間", "時長", "車種", "車號", "座位", "價錢",
                    "車票類型", "訂票狀態", "訂票代碼", "備註"]


def workbook():
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    back = wb.create_sheet("彰化回台北")
    back.append(PACK_HEADER)
    back.append(["重要", "錢包", None, None, None, "未打包", "未確認", "出發前", "需帶", "彰化", None])
    back.append([None, "鑰匙", "家鑰匙", None, None, "已打包", "確認", "隨時", "需帶", "彰化", None])
    back.append([None, None, "宿舍鑰匙", 1.0, None, "不需打包", "不需確認", "出發前晚", None, "彰化", "8/14沒帶"])
    back.append(["書 / 文具", "無"])
    there = wb.create_sheet("台北去彰化")
    there.append(PACK_HEADER)
    there.append(["食物", "水果", None, 2.0, 1.0, "未打包", "不需確認", "出發當天", "需買", "新北", None])
    there.append(["食物", "餅乾"])
    wb.create_sheet("Ignored tab").append(["anything"])
    transport = wb.create_sheet("Transportation")
    transport.append(TRANSPORT_HEADER)
    row = ["彰化火車站", "宿舍", "彰化客運6933A", "不需要", "http://map", "http://tt", None,
           "往鹿港乘車處", "高鐵台中", "鹿港", "彰化", "南瑤宮", 22.0, "7m", None,
           None, None, None, None,
           "7:30, *9:15", time(11, 35), "*14:15", 0.7777777777777778]
    transport.append(row)
    transport.append(["彰化火車站", "宿舍", "彰化客運6912"] + [None] * 16 + ["...", "...", None, None])
    transport.append(["宿舍", "彰化火車站"])
    transport.append(["新烏日火車站"])
    transport.append(["彰化火車站", "台北車站"] + [None] * 11 + ["2h-2h30m"])
    this_time = wb.create_sheet("This time")
    this_time.append(THIS_TIME_HEADER)
    this_time.append(["彰化火車站", "台北車站", "Thu 18:06-20:59", "2h53m", "火車 - 自強", 5158.0,
                      "5車15號", 550.0, "電子", "已訂票, 付款, 取票", 5891150.0, "提前買晚餐車上吃"])
    this_time.append(["台北車站", "彰化火車站", "Mon 12:15-14:23", "2h08m", "火車 - 自強 (3000)", 137.0,
                      "3車31號", 550.0, "電子", "已訂票, 付款", 5891246.0, "不吃午餐"])
    return wb
